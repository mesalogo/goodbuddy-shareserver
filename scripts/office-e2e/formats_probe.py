"""Real Docs 10 conversion/open/save probe. Requires libreoffice, Playwright, python-docx,
openpyxl and python-pptx. Uses isolated API/data and the existing Docs engine; never
restarts the running ShareServer. Build the backend before running this script.
"""
import argparse
import hashlib
import json
import os
import re
from pathlib import Path
import shutil
import subprocess
import tempfile
import time
import urllib.request
import zipfile
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from playwright.sync_api import sync_playwright

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--host', default='192.168.0.23')
parser.add_argument('--docs-port', type=int, default=8080, help='Isolated Docs engine port when testing a new image')
parser.add_argument('--env-file', default='.env.office')
parser.add_argument('--browser', help='Optional Chromium executable')
parser.add_argument('--desktop', action='store_true', help='Also run the adjacent GB legacy-DOC Electron save-as probe')
parser.add_argument('--desktop-only', action='store_true', help='Run only desktop normal and legacy save flows against the isolated stack')
parser.add_argument('--ai-only', action='store_true', help='Run three-format Electron right-click AI dialog probe')
parser.add_argument('--ai', action='store_true', help='Also verify three-format Electron right-click AI after browser checks')
parser.add_argument('--ai-formats', nargs='+', choices=['docx','xlsx','pptx'], default=['docx','xlsx','pptx'])
parser.add_argument('--chat-only', action='store_true', help='Run full built GB dialog and local Model Runtime probe')
parser.add_argument('--branding-screenshots', help='Save branded editor screenshots and check local branding requests')
parser.add_argument('--custom-ui', action='store_true', help='Validate opt-in GoodBuddy desktop branding and layout')
args = parser.parse_args()
repo = Path(__file__).resolve().parents[2]
root = Path(tempfile.mkdtemp(prefix='office-formats-', dir=repo / 'data'))
api = f'http://{args.host}:18087'
web = f'http://{args.host}:18088'
processes = []

def request(path, body=None):
    req = urllib.request.Request(api + '/api/v1/office' + path, data=body,
                                 headers={'Content-Type': 'application/octet-stream'})
    with urllib.request.urlopen(req, timeout=65) as response:
        return response.read()

def xml_contains(content, marker):
    import io
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        return any(marker in package.read(name).decode('utf-8', errors='replace')
                   for name in package.namelist() if name.endswith('.xml'))

try:
    doc = Document(); doc.add_paragraph('WORD_SOURCE_MARKER'); doc.save(root / 'sample.docx')
    book = Workbook(); book.active['A1'] = 'CELL_SOURCE_MARKER'; book.active['B1'] = '=1+2'; book.save(root / 'sample.xlsx')
    slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[0]); slide.shapes.title.text = 'SLIDE_SOURCE_MARKER'; slides.save(root / 'sample.pptx')
    for modern, legacy, conversion in [('docx', 'doc', 'doc:MS Word 97'), ('xlsx', 'xls', 'xls:MS Excel 97'), ('pptx', 'ppt', 'ppt:MS PowerPoint 97')]:
        subprocess.run(['libreoffice', f'-env:UserInstallation={(root / "lo-profile").as_uri()}', '--headless',
                        '--convert-to', conversion, '--outdir', str(root), str(root / f'sample.{modern}')], check=True, timeout=60)
        assert (root / f'sample.{legacy}').is_file()
    env = dict(os.environ, OFFICE_ENABLED='true', OFFICE_DATA_DIR=str(root / 'office'), SHARESERVER_DATA_DIR=str(root / 'metadata'),
               SHARESERVER_HOST='0.0.0.0', SHARESERVER_PORT='18087', SHARESERVER_API_PORT='18087', SHARESERVER_PUBLIC_URL=web, DATABASE_URL='',
               OFFICE_INTERNAL_URL=api, ONLYOFFICE_PUBLIC_URL=f'http://{args.host}:{args.docs_port}', ONLYOFFICE_INTERNAL_URL=f'http://{args.host}:{args.docs_port}')
    log = (root / 'services.log').open('w')
    if args.custom_ui:
        env['OFFICE_CUSTOM_UI_ENABLED'] = 'true'
        env['OFFICE_WHITE_LABEL_ENABLED'] = 'false'
    processes.append(subprocess.Popen(['node', '--env-file=' + str(repo / args.env_file), 'dist/server/server/src/index.js'], cwd=repo, env=env, stdout=log, stderr=log))
    processes.append(subprocess.Popen(['node', 'node_modules/vite/bin/vite.js', '--config', 'web/vite.config.ts', '--host', '0.0.0.0', '--port', '18088'], cwd=repo, env=env, stdout=log, stderr=log))
    for attempt in range(60):
        try:
            assert json.loads(request('/config'))['enabled']
            with urllib.request.urlopen(web, timeout=2) as response: assert response.status == 200
            break
        except Exception:
            if any(p.poll() is not None for p in processes): raise RuntimeError('Probe service exited')
            time.sleep(.25)
    else: raise RuntimeError('Probe services unavailable')
    records = []
    for extension, marker in [('docx', 'WORD_SOURCE_MARKER'), ('xlsx', 'CELL_SOURCE_MARKER'), ('pptx', 'SLIDE_SOURCE_MARKER'),
                              ('doc', 'WORD_SOURCE_MARKER'), ('xls', 'CELL_SOURCE_MARKER'), ('ppt', 'SLIDE_SOURCE_MARKER')]:
        source = (root / f'sample.{extension}').read_bytes()
        record = json.loads(request('/documents?name=sample.' + extension, source))
        output = request(f'/documents/{record["id"]}/download?revision=1')
        assert xml_contains(output, marker), extension
        assert record['digest'] == hashlib.sha256(output).hexdigest()
        assert record['requiresSaveAs'] == (extension in ('doc', 'xls', 'ppt'))
        records.append((extension, record))
        print(f'PASS: {extension} upload/conversion -> {record["name"]}; marker and digest verified', flush=True)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(executable_path=args.browser, headless=True, args=['--no-sandbox'])
        for extension, record in ([] if args.desktop_only or args.ai_only or args.chat_only else records):
            page = browser.new_page(viewport={'width': 1440, 'height': 1000})
            branding_responses = []
            external_requests = set()
            page.on('response', lambda response: branding_responses.append((response.url, response.status)) if '/api/v1/office/branding/' in response.url else None)
            page.on('request', lambda request: external_requests.add(request.url.split('/')[2]) if request.url.startswith(('http://','https://')) and request.url.split('/')[2].split(':')[0] not in (args.host, 'localhost', '127.0.0.1') else None)
            page.goto(web + '/office/editor?document=' + record['id'])
            page.wait_for_function('document.documentElement.dataset.officeReady === "true"', timeout=60000)
            frame = page.frames[1]
            if frame.get_by_role('button', name='取消', exact=True).count():
                frame.get_by_role('button', name='取消', exact=True).click()
            print(f'PASS: {extension} real browser editor ready', flush=True)
            if args.custom_ui:
                from urllib.parse import parse_qs, urlparse
                params = parse_qs(urlparse(frame.url).query)
                assert params.get('customer') == ['GoodBuddy · Powered by ONLYOFFICE'], params
                assert '/api/v1/office/branding/icon-light.png' in params.get('logo', [''])[0]
                assert any('/logo-' in url and status == 200 for url, status in branding_responses)
                for label in ['插件', 'Plugins']:
                    assert not any(item.is_visible() for item in frame.get_by_text(label, exact=True).all()), label
                print('PASS:', extension, 'custom loader parameters, local header logo and hidden Plugins tab', flush=True)
            if args.branding_screenshots:
                page.wait_for_function('document.documentElement.dataset.officeReady === "true"')
                page.screenshot(path=str(Path(args.branding_screenshots) / f'goodbuddy-office-{extension}.png'))
                assert any('/logo-' in url and status == 200 for url,status in branding_responses), branding_responses
                assert frame.get_by_role('button', name=re.compile(r'^(Chat|聊天)(\s|$)')).count() == 0
                print('BRANDING:', extension, 'local logo loaded, chat hidden; external browser hosts:', sorted(external_requests), flush=True)
            # Edit via the actual editor keyboard; inspect persisted OOXML, not just UI state.
            modern = record['name'].split('.')[-1]
            marker = 'GB_FORMAT_EDIT_' + extension.upper()
            if modern == 'docx':
                page.mouse.click(650, 560); page.keyboard.press('Control+End'); page.keyboard.type(marker)
            elif modern == 'xlsx':
                canvas = frame.locator('#ws-canvas')
                box = canvas.bounding_box()
                if not box: raise RuntimeError('Spreadsheet canvas not found')
                page.mouse.click(box['x'] + 90, box['y'] + 40)
                page.keyboard.press('Control+Home'); page.keyboard.type(marker); page.keyboard.press('Enter')
            else:
                page.screenshot(path=str(root / f'{extension}-editor.png'))
                # Title slide: enter the existing title shape and replace only its text.
                page.mouse.dblclick(800, 475)
                page.keyboard.press('Control+a'); page.keyboard.type(marker); page.keyboard.press('Escape')
            page.keyboard.press('Control+s')
            page.wait_for_function('document.documentElement.dataset.officeDirty === "false"', timeout=30000)
            # Native editor Save must persist without the removed host-page save button.
            for attempt in range(30):
                page.wait_for_timeout(500)
                output = request('/documents/' + record['id'] + '/download')
                if xml_contains(output, marker): break
            else:
                page.screenshot(path=str(root / f'{extension}-failed.png'))
                raise RuntimeError(f'{extension}: edited marker missing from saved file')
            print(f'PASS: {extension} keyboard edit -> real callback -> persisted output', flush=True)
            if args.branding_screenshots:
                frame.get_by_text('文件', exact=True).first.click()
                for label in ['帮助', '反馈与支持', 'Help', 'Feedback & Support', 'Suggest a Feature']:
                    assert not any(item.is_visible() for item in frame.get_by_text(label, exact=True).all()), label
                if extension == 'docx':
                    page.screenshot(path=str(Path(args.branding_screenshots) / 'goodbuddy-office-file-menu.png'))
                print('PASS:', extension, 'File menu hides help/feedback; external browser hosts:', sorted(external_requests), flush=True)
            page.close()
        browser.close()
    if args.desktop or args.desktop_only:
        import sys
        if args.desktop_only:
            subprocess.run([sys.executable, str(repo.parent / 'goodbuddy/scripts/verify-office-desktop.py'), '--web', web, '--file-entry'], cwd=repo.parent / 'goodbuddy', check=True, timeout=150)
        subprocess.run([sys.executable, str(repo.parent / 'goodbuddy/scripts/verify-office-desktop.py'), '--web', web, '--legacy'], cwd=repo.parent / 'goodbuddy', check=True, timeout=120)
    if args.ai_only or args.ai:
        import sys
        subprocess.run([sys.executable, str(repo.parent / 'goodbuddy/scripts/verify-office-ai.py'), '--web', web, '--formats', *args.ai_formats], cwd=repo.parent / 'goodbuddy', check=True, timeout=360)
    if args.chat_only:
        import sys
        subprocess.run([sys.executable, str(repo.parent / 'goodbuddy/scripts/verify-office-chat.py'), '--web', web], cwd=repo.parent / 'goodbuddy', check=True, timeout=240)
    print('PASS: selected format/desktop flows; isolated probe data will be removed', flush=True)
finally:
    for process in processes:
        if process.poll() is None:
            process.terminate()
            try: process.wait(timeout=10)
            except subprocess.TimeoutExpired: process.kill(); process.wait()
    import sys
    if sys.exc_info()[0]: print('Probe diagnostics retained at', root, flush=True)
    else: shutil.rmtree(root)
